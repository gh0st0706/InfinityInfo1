
import asyncio
import base64
import csv
import html.parser
import ipaddress
import io
import json
import logging
import os
import re
import socket
import sys
import tempfile
import time
import xml.etree.ElementTree as ET
from pathlib import Path
from urllib.parse import parse_qs, urljoin, urlparse
from typing import Any

import dns.asyncresolver
import phonenumbers

import httpx
from dotenv import dotenv_values, load_dotenv
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    Workbook = load_workbook = None

try:
    import xlrd
except ImportError:
    xlrd = None

load_dotenv()
logger = logging.getLogger(__name__)


def load_hunter_api_key():
    key = os.getenv("HUNTER_API_KEY", "").strip()
    if key:
        return key

    app_directory = Path(__file__).resolve().parent
    for key_file in (
        Path("/etc/secrets/HUNTER_API_KEY"),
        app_directory / "HUNTER_API_KEY",
    ):
        try:
            key = key_file.read_text(encoding="utf-8").strip()
        except OSError:
            continue
        if key:
            return key

    for env_file in (
        Path("/etc/secrets/.env"),
        app_directory / ".env",
    ):
        key = (dotenv_values(env_file).get("HUNTER_API_KEY") or "").strip()
        if key:
            return key

    return ""

app = FastAPI(title="InfinityInfo API")
allowed_origins = [
    "https://infinityinfo1-infinityinfo.vercel.app",
    *[
        origin.strip()
        for origin in os.getenv(
            "CORS_ALLOWED_ORIGINS",
            "http://localhost:5500,http://127.0.0.1:5500,null",
        ).split(",")
        if origin.strip()
    ],
]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed_origins,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
    expose_headers=["Content-Disposition"],
)

HUNTER_API_KEY = load_hunter_api_key()
HUNTER_URL = "https://api.hunter.io/v2/email-finder"
TWILIO_LOOKUP_URL = "https://lookups.twilio.com/v2/PhoneNumbers/"
TWILIO_API_KEY = os.getenv("TWILIO_API_KEY", "").strip()
TWILIO_API_SECRET = os.getenv("TWILIO_API_SECRET", "").strip()
DEFAULT_PHONE_COUNTRY = os.getenv("DEFAULT_COUNTRY", "IN").strip().upper()
MAX_PHONE_BULK_ROWS = max(1, int(os.getenv("MAX_BULK_ROWS", "500")))
MAX_UPLOAD_BYTES = 4 * 1024 * 1024

REQUIRED_COLUMNS = {
    "First Name",
    "Last Name",
    "Full Name",
    "Company Website",
    "Email",
    "Email Status",
}


class EmailSearch(BaseModel):
    first_name: str
    last_name: str
    domain: str


class PersonFinderSearch(BaseModel):
    full_name: str
    company: str


class PhoneValidationRequest(BaseModel):
    phone: str
    country: str | None = None
    useProvider: bool = False


class PhoneBulkValidationRequest(BaseModel):
    numbers: list[Any]
    country: str | None = None
    useProvider: bool = False


class PublicPageParser(html.parser.HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.text_parts = []
        self.links = []

    def handle_data(self, data):
        self.text_parts.append(data)

    def handle_starttag(self, tag, attrs):
        if tag.lower() == "a":
            href = dict(attrs).get("href", "")
            if href:
                self.links.append(href)
                if href.lower().startswith("mailto:"):
                    self.text_parts.append(href[7:].split("?", 1)[0])
                elif href.lower().startswith("tel:"):
                    self.text_parts.append(href[4:].split("?", 1)[0])


EMAIL_PATTERN = re.compile(r"[A-Z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?(?:\.[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?)+", re.I)


def normalize_company_domain(value):
    value = clean(value)
    parsed = urlparse(value if "://" in value else f"https://{value}")
    host = (parsed.hostname or "").lower().rstrip(".")
    if host.startswith("www."):
        host = host[4:]
    if not host or "." not in host or len(host) > 253:
        raise HTTPException(status_code=400, detail="Enter a valid company domain")
    try:
        ipaddress.ip_address(host)
    except ValueError:
        pass
    else:
        raise HTTPException(status_code=400, detail="Enter a company domain, not an IP address")
    if any(not label or len(label) > 63 or not re.fullmatch(r"[a-z0-9](?:[a-z0-9-]*[a-z0-9])?", label) for label in host.split(".")):
        raise HTTPException(status_code=400, detail="Enter a valid company domain")
    return host


async def assert_public_host(host):
    try:
        addresses = await asyncio.to_thread(
            socket.getaddrinfo, host, 443, type=socket.SOCK_STREAM
        )
    except OSError as exc:
        raise HTTPException(status_code=400, detail="The company domain could not be resolved") from exc
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise HTTPException(status_code=400, detail="The company domain does not resolve to a public address")


async def inspect_public_company_site(client, domain):
    await assert_public_host(domain)
    allowed_hosts = {domain}
    urls = [f"https://{domain}/"]
    try:
        await assert_public_host(f"www.{domain}")
    except HTTPException:
        pass
    else:
        allowed_hosts.add(f"www.{domain}")
        urls.append(f"https://www.{domain}/")

    # A lot of company sites keep their contact information off the homepage.
    # Seed the common public pages so the finder does not depend on navigation
    # order or a site's home page being rendered as static HTML.
    urls.extend(
        f"https://{domain}/{path}"
        for path in ("contact", "contact-us", "about", "about-us", "team", "leadership")
    )
    visited = set()
    checked_pages = set()
    found = {}
    pages_checked = 0
    deadline = asyncio.get_running_loop().time() + 5

    async def read_public_page(url, require_html=True):
        for _ in range(6):
            parsed = urlparse(url)
            if (
                parsed.scheme not in {"http", "https"}
                or parsed.hostname not in allowed_hosts
                or parsed.port not in {None, 80, 443}
                or parsed.username
                or parsed.password
            ):
                return None
            try:
                async with client.stream(
                    "GET", url, headers={"User-Agent": "InfinityInfoEmailFinder/1.0"},
                    follow_redirects=False,
                ) as response:
                    if response.is_redirect:
                        location = response.headers.get("location")
                        if not location:
                            return None
                        url = urljoin(url, location)
                        continue
                    content_type = response.headers.get("content-type", "").lower()
                    if response.status_code >= 400 or (require_html and "html" not in content_type):
                        return None
                    body = bytearray()
                    async for chunk in response.aiter_bytes():
                        body.extend(chunk)
                        if len(body) > 1_000_000:
                            break
                    return bytes(body), url
            except (httpx.HTTPError, ValueError):
                return None
        return None

    # Use the site's own robots.txt sitemap declarations and conventional
    # sitemap locations to discover public contact and staff pages.
    sitemap_urls = [f"https://{domain}/sitemap.xml"]
    try:
        robots = await asyncio.wait_for(
            read_public_page(f"https://{domain}/robots.txt", require_html=False),
            timeout=min(0.4, max(0.1, deadline - asyncio.get_running_loop().time())),
        )
    except asyncio.TimeoutError:
        robots = None
    if robots:
        robots_text = robots[0].decode("utf-8", errors="ignore")
        sitemap_urls.extend(
            line.split(":", 1)[1].strip()
            for line in robots_text.splitlines()
            if line.lower().startswith("sitemap:") and ":" in line
        )

    discovered_urls = []
    sitemap_queue = list(dict.fromkeys(sitemap_urls))
    sitemap_seen = set()
    for _ in range(3):
        if not sitemap_queue:
            break
        sitemap_url = sitemap_queue.pop(0)
        if sitemap_url in sitemap_seen:
            continue
        sitemap_seen.add(sitemap_url)
        remaining = deadline - asyncio.get_running_loop().time()
        if remaining <= 0:
            break
        try:
            sitemap = await asyncio.wait_for(
                read_public_page(sitemap_url, require_html=False), timeout=min(0.4, remaining)
            )
        except asyncio.TimeoutError:
            break
        if not sitemap:
            continue
        try:
            root = ET.fromstring(sitemap[0])
        except ET.ParseError:
            continue
        is_sitemap_index = root.tag.rsplit("}", 1)[-1] == "sitemapindex"
        for element in root.iter():
            if element.tag.rsplit("}", 1)[-1] != "loc" or not element.text:
                continue
            candidate_url = element.text.strip()
            parsed_candidate = urlparse(candidate_url)
            if (
                is_sitemap_index
                and parsed_candidate.scheme in {"http", "https"}
                and parsed_candidate.hostname in allowed_hosts
                and parsed_candidate.port in {None, 80, 443}
                and "sitemap" in parsed_candidate.path.lower()
            ):
                sitemap_queue.append(candidate_url)
                continue
            if (
                parsed_candidate.scheme in {"http", "https"}
                and parsed_candidate.hostname in allowed_hosts
                and parsed_candidate.port in {None, 80, 443}
                and not parsed_candidate.username
                and not parsed_candidate.password
                and re.search(r"contact|about|team|leadership|people|staff|press|media", parsed_candidate.path, re.I)
            ):
                discovered_urls.append(candidate_url.split("#", 1)[0])
                if len(discovered_urls) >= 12:
                    break
        if discovered_urls:
            break
    urls.extend(discovered_urls)

    while urls and pages_checked < 12:
        url = urls.pop(0)
        parsed_url = urlparse(url)
        normalized_url = url.split("#", 1)[0]
        if parsed_url.hostname not in allowed_hosts or normalized_url in visited:
            continue
        visited.add(normalized_url)
        remaining = deadline - asyncio.get_running_loop().time()
        if remaining <= 0:
            break
        try:
            page = await asyncio.wait_for(read_public_page(url), timeout=remaining)
        except asyncio.TimeoutError:
            break
        if not page:
            continue
        pages_checked += 1
        body, final_url = page
        final_url = final_url.split("#", 1)[0]
        visited.add(final_url)
        checked_pages.add(final_url)
        parser = PublicPageParser()
        parser.feed(body.decode("utf-8", errors="ignore"))
        page_text = " ".join(parser.text_parts)
        page_text = re.sub(r"\s*(?:\[at\]|\(at\)|\{at\})\s*", "@", page_text, flags=re.I)
        page_text = re.sub(r"\s*(?:\[dot\]|\(dot\)|\{dot\})\s*", ".", page_text, flags=re.I)
        for email in EMAIL_PATTERN.findall(page_text):
            normalized_email = email.lower()
            found.setdefault(normalized_email, set()).add(final_url)
        for href in parser.links:
            next_url = urljoin(final_url, href).split("#", 1)[0]
            next_parsed = urlparse(next_url)
            if (
                next_parsed.scheme in {"http", "https"}
                and next_parsed.hostname in allowed_hosts
                and next_parsed.port in {None, 80, 443}
                and not next_parsed.username
                and not next_parsed.password
                and next_url not in visited
            ):
                urls.append(next_url)
    return {
        "emails": {
            email: sorted(sources)
            for email, sources in found.items()
            if email.rsplit("@", 1)[1] == domain
        },
        "pages_checked": sorted(checked_pages),
    }


async def lookup_person_publicly(first_name, last_name, domain):
    resolver = dns.asyncresolver.Resolver()
    try:
        mx_records = await resolver.resolve(domain, "MX", lifetime=4)
        if all(record.exchange.to_text() == "." for record in mx_records):
            return {"found": False, "email": "", "confidence": "", "validation": "", "result": "no_mail_server", "candidates": [], "sources": [], "pages_checked": []}
    except dns.resolver.NXDOMAIN:
        return {"found": False, "email": "", "confidence": "", "validation": "", "result": "no_mail_server", "candidates": [], "sources": [], "pages_checked": []}
    except dns.resolver.NoAnswer:
        # RFC 5321 permits an implicit MX: the domain itself accepts mail.
        try:
            await resolver.resolve(domain, "A", lifetime=3)
        except dns.resolver.NoAnswer:
            try:
                await resolver.resolve(domain, "AAAA", lifetime=3)
            except (dns.resolver.NoAnswer, dns.resolver.NXDOMAIN):
                return {"found": False, "email": "", "confidence": "", "validation": "", "result": "no_mail_server", "candidates": [], "sources": [], "pages_checked": []}
    except (dns.resolver.NoNameservers, dns.exception.Timeout, dns.exception.DNSException) as exc:
        raise HTTPException(status_code=503, detail="Could not check company mail records. Please try again.") from exc

    async with httpx.AsyncClient(timeout=8, follow_redirects=False) as client:
        public_site = await inspect_public_company_site(client, domain)
    public_evidence = public_site["emails"]
    pages_checked = public_site["pages_checked"]

    first = first_name.casefold()
    last = last_name.casefold()
    direct_matches = [email for email in public_evidence if email.split("@", 1)[0].replace(".", "").replace("_", "").replace("-", "") in {first + last, last + first}]
    if direct_matches:
        sources = sorted({source for email in direct_matches for source in public_evidence[email]})
        return {"found": True, "email": direct_matches[0], "confidence": "public_source", "validation": "domain_has_mx", "result": "found_publicly", "candidates": [], "sources": sources[:5], "pages_checked": pages_checked[:8]}

    separators = [".", "_", "-"]
    generic_mailboxes = {
        "abuse", "admin", "billing", "careers", "contact", "hello", "hr", "info",
        "jobs", "legal", "media", "noreply", "no-reply", "office", "postmaster",
        "press", "privacy", "sales", "security", "support", "team", "webmaster",
    }
    sample_locals = []
    for email in public_evidence:
        local = email.split("@", 1)[0].lower()
        first_local_token = re.split(r"[._+-]", local, maxsplit=1)[0]
        if local not in generic_mailboxes and first_local_token not in generic_mailboxes:
            sample_locals.append(local)
    separator_samples = {
        separator: [
            local for local in sample_locals
            if re.fullmatch(r"[a-z]+" + re.escape(separator) + r"[a-z]+", local)
        ]
        for separator in separators
    }
    observed_separator = next(
        (separator for separator in separators if separator_samples[separator]), None
    )
    if observed_separator:
        candidate = f"{first}{observed_separator}{last}@{domain}"
        confidence = "domain_format_clue" if len(separator_samples[observed_separator]) >= 2 else "low"
        result = "pattern_candidate"
        candidates = [candidate]
    elif sample_locals:
        candidate = f"{first}{last}@{domain}"
        confidence = "low"
        result = "pattern_candidate"
        candidates = [candidate]
    else:
        candidate = ""
        confidence = ""
        result = "not_found_publicly"
        candidates = []

    sources = sorted({source for found_sources in public_evidence.values() for source in found_sources})
    return {"found": False, "email": candidate, "confidence": confidence, "validation": "domain_has_mx", "result": result, "candidates": candidates, "sources": sources[:5], "pages_checked": pages_checked[:8]}


def validate_phone_locally(value, country=None):
    raw = clean(value)
    region = clean(country).upper() or DEFAULT_PHONE_COUNTRY
    if region not in phonenumbers.SUPPORTED_REGIONS:
        region = DEFAULT_PHONE_COUNTRY if DEFAULT_PHONE_COUNTRY in phonenumbers.SUPPORTED_REGIONS else "IN"
    result = {
        "input": raw,
        "normalized": "",
        "nationalFormat": "",
        "country": "",
        "callingCode": "",
        "validFormat": False,
        "possible": False,
        "lineType": "unknown",
        "carrier": "",
        "lookupStatus": "local_only",
        "reachability": "not_checked",
        "whatsapp": "not_checked",
        "providerConfigured": bool(TWILIO_API_KEY and TWILIO_API_SECRET),
    }
    if not raw:
        result["message"] = "Empty phone number"
        return result
    try:
        parsed = phonenumbers.parse(raw, region)
    except phonenumbers.NumberParseException:
        result["message"] = "Could not parse number; include a country code such as +91"
        return result

    valid = phonenumbers.is_valid_number(parsed)
    possible = phonenumbers.is_possible_number(parsed)
    type_labels = {
        phonenumbers.PhoneNumberType.FIXED_LINE: "fixed_line",
        phonenumbers.PhoneNumberType.MOBILE: "mobile",
        phonenumbers.PhoneNumberType.FIXED_LINE_OR_MOBILE: "fixed_line_or_mobile",
        phonenumbers.PhoneNumberType.TOLL_FREE: "toll_free",
        phonenumbers.PhoneNumberType.PREMIUM_RATE: "premium_rate",
        phonenumbers.PhoneNumberType.SHARED_COST: "shared_cost",
        phonenumbers.PhoneNumberType.VOIP: "voip",
        phonenumbers.PhoneNumberType.PERSONAL_NUMBER: "personal_number",
        phonenumbers.PhoneNumberType.PAGER: "pager",
        phonenumbers.PhoneNumberType.UAN: "uan",
        phonenumbers.PhoneNumberType.VOICEMAIL: "voicemail",
    }
    result.update({
        "normalized": phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.E164),
        "nationalFormat": phonenumbers.format_number(parsed, phonenumbers.PhoneNumberFormat.NATIONAL),
        "country": phonenumbers.region_code_for_number(parsed) or "",
        "callingCode": str(parsed.country_code or ""),
        "validFormat": valid,
        "possible": possible,
        "lineType": type_labels.get(phonenumbers.number_type(parsed), "unknown"),
        "message": "Matches numbering-plan metadata; active service is not confirmed" if valid else "Does not match current numbering-plan metadata",
    })
    return result


async def enrich_phone_with_provider(result):
    if not result.get("normalized") or not result.get("validFormat"):
        return result
    credentials = (TWILIO_API_KEY, TWILIO_API_SECRET)
    try:
        async with httpx.AsyncClient(timeout=15) as client:
            response = await client.get(
                f"{TWILIO_LOOKUP_URL}{result['normalized']}",
                params={"Fields": "line_type_intelligence"},
                auth=credentials,
                headers={"Accept": "application/json"},
            )
        if response.status_code >= 400:
            raise httpx.HTTPStatusError("Lookup provider returned an error", request=response.request, response=response)
        data = response.json()
        line_info = data.get("line_type_intelligence") or {}
        result.update({
            "providerValid": data.get("valid"),
            "normalized": data.get("phone_number") or result["normalized"],
            "nationalFormat": data.get("national_format") or result["nationalFormat"],
            "country": data.get("country_code") or result["country"],
            "callingCode": data.get("calling_country_code") or result["callingCode"],
            "carrier": line_info.get("carrier_name") or "",
            "lineType": line_info.get("type") or result["lineType"],
            "lookupStatus": "provider_checked",
            "reachability": "not_confirmed",
            "providerErrorCode": line_info.get("error_code"),
            "message": "Provider lookup completed. Numbering-plan validity does not guarantee an active or reachable line.",
        })
        if data.get("valid") is not None:
            result["validFormat"] = data["valid"]
        return result
    except (httpx.HTTPError, ValueError):
        result.update({
            "lookupStatus": "provider_error",
            "message": result["message"] + "; carrier lookup did not complete.",
        })
        return result


async def check_phone(value, country=None, use_provider=False):
    result = validate_phone_locally(value, country)
    if use_provider and TWILIO_API_KEY and TWILIO_API_SECRET:
        return await enrich_phone_with_provider(result)
    return result


def clean(value):
    return str(value).strip() if value is not None else ""


def get_domain(website):
    website = clean(website)
    if not website:
        return ""

    parsed = urlparse(
        website if "://" in website else f"https://{website}"
    )
    host = (parsed.hostname or "").lower()

    if host.startswith("www."):
        host = host[4:]

    return host


def split_full_name(full_name):
    parts = clean(full_name).split()
    if len(parts) < 2:
        return "", ""

    return parts[0], " ".join(parts[1:])


async def find_with_hunter(client, first_name, last_name, domain):
    if not HUNTER_API_KEY:
        raise RuntimeError("Hunter API key is not configured")

    response = await client.get(
        HUNTER_URL,
        params={
            "first_name": first_name,
            "last_name": last_name,
            "domain": domain,
            "api_key": HUNTER_API_KEY,
        },
    )

    if response.status_code != 200:
        # Do not return raw provider errors or secrets to the user.
        return {
            "email": "",
            "confidence": "",
            "verification": "",
            "result": f"provider_error_{response.status_code}",
        }

    data = response.json().get("data", {})
    email = clean(data.get("email"))

    verification = data.get("verification") or {}
    status = clean(verification.get("status"))

    return {
        "email": email,
        "confidence": data.get("score", ""),
        "verification": status,
        "result": "found" if email else "not_found",
    }


@app.get("/health")
@app.get("/api/health")
def home():
    return {
        "ok": True,
        "message": "InfinityInfo API is running",
        "hunter_api_configured": bool(HUNTER_API_KEY),
        "public_lookup_available": True,
        "providerConfigured": bool(TWILIO_API_KEY and TWILIO_API_SECRET),
        "defaultCountry": DEFAULT_PHONE_COUNTRY,
        "maxBulkRows": MAX_PHONE_BULK_ROWS,
    }


# Keep the existing Render health check working during the migration. Vercel
# serves public/index.html at /, so the API root is intentionally omitted there.
if not os.getenv("VERCEL"):
    app.add_api_route(
        "/",
        home,
        methods=["GET"],
        include_in_schema=False,
    )


@app.post("/find-email")
@app.post("/api/find-email")
async def find_email(search: EmailSearch):
    if not HUNTER_API_KEY:
        raise HTTPException(
            status_code=500,
            detail="Hunter API key is not configured",
        )

    async with httpx.AsyncClient(timeout=30) as client:
        result = await find_with_hunter(
            client,
            search.first_name,
            search.last_name,
            search.domain,
        )

    return {
        "name": f"{search.first_name} {search.last_name}",
        **result,
    }


@app.post("/email-finder")
@app.post("/api/email-finder")
async def email_finder(search: PersonFinderSearch):
    name_parts = clean(search.full_name).split()
    if len(name_parts) < 2:
        raise HTTPException(status_code=400, detail="Enter a first and last name")
    if len(search.full_name) > 120:
        raise HTTPException(status_code=400, detail="Name must be 120 characters or fewer")
    domain = normalize_company_domain(search.company)
    result = await lookup_person_publicly(name_parts[0], name_parts[-1], domain)
    return {"name": clean(search.full_name), "domain": domain, **result}


class PublicSearchResultParser(html.parser.HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.results = []
        self.current_title = None
        self.current_url = ""
        self.snippet_depth = 0
        self.current_snippet = []
        self.snippet_target = None

    def handle_starttag(self, tag, attrs):
        attributes = dict(attrs)
        classes = set((attributes.get("class") or "").split())
        if tag == "a" and "result__a" in classes:
            self.current_title = []
            self.current_url = attributes.get("href", "")
        if "result__snippet" in classes and self.results:
            self.snippet_target = self.results[-1]
            self.current_snippet = []
            self.snippet_depth = 1
        elif self.snippet_depth:
            self.snippet_depth += 1

    def handle_data(self, data):
        if self.current_title is not None:
            self.current_title.append(data)
        if self.snippet_depth:
            self.current_snippet.append(data)

    def handle_endtag(self, tag):
        if tag == "a" and self.current_title is not None:
            self.results.append({
                "title": " ".join("".join(self.current_title).split()),
                "url": self.current_url,
                "snippet": "",
            })
            self.current_title = None
            self.current_url = ""
        elif self.snippet_depth:
            self.snippet_depth -= 1
            if self.snippet_depth == 0 and self.snippet_target is not None:
                self.snippet_target["snippet"] = " ".join(" ".join(self.current_snippet).split())
                self.snippet_target = None
                self.current_snippet = []


COMPANY_SEARCH_CACHE = {}
COMPANY_SEARCH_CACHE_TTL_SECONDS = 1800


def public_search_result_url(href):
    parsed = urlparse(urljoin("https://duckduckgo.com/", href))
    if parsed.hostname and (parsed.hostname == "duckduckgo.com" or parsed.hostname.endswith(".duckduckgo.com")):
        target = parse_qs(parsed.query).get("uddg", [""])[0]
        if target:
            parsed = urlparse(target)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        return ""
    if parsed.hostname == "duckduckgo.com" or parsed.hostname.endswith(".duckduckgo.com"):
        return ""
    try:
        port = parsed.port
    except ValueError:
        return ""
    if port not in {None, 80, 443}:
        return ""
    return parsed.geturl()


async def read_search_result_page(client, result, location, country):
    source_url = public_search_result_url(result.get("url", ""))
    if not source_url:
        return None
    parsed_source = urlparse(source_url)
    try:
        await assert_public_host(parsed_source.hostname)
    except HTTPException:
        return None

    page_url = source_url
    try:
        async with asyncio.timeout(4):
            async with client.stream(
                "GET", page_url,
                headers={"User-Agent": "InfinityInfoPublicSearch/1.0 (contact: info@infinityinfo.com)"},
                follow_redirects=False,
            ) as response:
                if response.is_redirect:
                    destination = urljoin(page_url, response.headers.get("location", ""))
                    parsed_destination = urlparse(destination)
                    if parsed_destination.hostname != parsed_source.hostname:
                        return None
                    response = await client.get(destination, follow_redirects=False)
                    if response.status_code >= 400:
                        return None
                    content_type = response.headers.get("content-type", "")
                    body = response.content[:500_000]
                    page_url = destination
                else:
                    if response.status_code >= 400 or "html" not in response.headers.get("content-type", "").lower():
                        return None
                    content_type = response.headers.get("content-type", "")
                    body = bytearray()
                    async for chunk in response.aiter_bytes():
                        body.extend(chunk)
                        if len(body) >= 500_000:
                            break
                    body = bytes(body)
        parser = PublicPageParser()
        parser.feed(body.decode("utf-8", errors="ignore"))
    except (httpx.HTTPError, asyncio.TimeoutError, ValueError):
        parser = PublicPageParser()
        parser.feed(result.get("snippet", ""))
        page_url = source_url

    text = " ".join(parser.text_parts)
    for href in parser.links:
        contact_url = urljoin(page_url, href)
        parsed_contact = urlparse(contact_url)
        try:
            contact_port = parsed_contact.port
        except ValueError:
            continue
        if (
            parsed_contact.scheme not in {"http", "https"}
            or parsed_contact.hostname != parsed_source.hostname
            or parsed_contact.username
            or parsed_contact.password
            or contact_port not in {None, 80, 443}
            or not re.search(r"contact|about|team|enquir|reach-us", parsed_contact.path, re.I)
        ):
            continue
        try:
            async with asyncio.timeout(2):
                async with client.stream(
                    "GET", contact_url,
                    headers={"User-Agent": "InfinityInfoPublicSearch/1.0"},
                    follow_redirects=False,
                ) as contact_response:
                    if contact_response.status_code >= 400 or "html" not in contact_response.headers.get("content-type", "").lower():
                        continue
                    contact_body = bytearray()
                    async for chunk in contact_response.aiter_bytes():
                        contact_body.extend(chunk)
                        if len(contact_body) >= 300_000:
                            break
            contact_parser = PublicPageParser()
            contact_parser.feed(bytes(contact_body).decode("utf-8", errors="ignore"))
            parser.text_parts.extend(contact_parser.text_parts)
            page_url = contact_url
            break
        except (httpx.HTTPError, asyncio.TimeoutError, ValueError):
            continue

    text = " ".join(parser.text_parts)
    emails = sorted({email.lower() for email in EMAIL_PATTERN.findall(text)})
    phone_values = []
    for candidate in re.findall(r"(?<!\d)(?:\+?\d[\d\s().-]{7,}\d)(?!\d)", text):
        digits = re.sub(r"\D", "", candidate)
        if 9 <= len(digits) <= 15:
            phone_values.append(re.sub(r"\s+", " ", candidate).strip())

    website = f"{urlparse(source_url).scheme}://{urlparse(source_url).netloc}"
    return {
        "companyName": result.get("title", "").strip()[:240],
        "address": "",
        "street": "",
        "city": location,
        "state": "",
        "postalCode": "",
        "country": country,
        "email": emails[0] if emails else "",
        "phone": phone_values[0] if phone_values else "",
        "website": website,
        "locationType": "Unverified",
        "source": "Public web search",
        "sourceUrl": page_url,
    }


@app.get("/search")
@app.get("/api/search")
async def search_companies(query: str, location: str = "", country: str = ""):
    query = clean(query)
    location = clean(location)
    country = clean(country)
    if not query:
        raise HTTPException(status_code=400, detail="Enter a company name, industry, or search phrase.")
    if len(query) > 180 or len(location) > 120 or len(country) > 80:
        raise HTTPException(status_code=400, detail="Search fields exceed the allowed length.")

    cache_key = "|".join((query.casefold(), location.casefold(), country.casefold()))
    cached = COMPANY_SEARCH_CACHE.get(cache_key)
    if cached and time.time() - cached[0] < COMPANY_SEARCH_CACHE_TTL_SECONDS:
        return cached[1]

    backend_directory = Path(__file__).resolve().parent
    spider_path = backend_directory / "company_spider.py"
    try:
        with tempfile.TemporaryDirectory(prefix="infinityinfo-spider-") as temp_dir:
            output_path = str(Path(temp_dir) / "results.json")
            spider_env = os.environ.copy()
            spider_env["INFINITYINFO_SPIDER_OUTPUT"] = output_path
            process = await asyncio.create_subprocess_exec(
                sys.executable, str(spider_path), query, location, country,
                cwd=str(backend_directory), env=spider_env,
                stdout=asyncio.subprocess.DEVNULL,
                stderr=asyncio.subprocess.PIPE,
            )
            try:
                _, stderr = await asyncio.wait_for(process.communicate(), timeout=19)
            except asyncio.TimeoutError:
                process.kill()
                await process.wait()
                raise
            if process.returncode != 0:
                detail = stderr.decode("utf-8", errors="replace")[-1200:]
                raise RuntimeError(detail or "The public web crawler could not start.")
            results_path = Path(output_path)
            enriched = json.loads(results_path.read_text(encoding="utf-8")) if results_path.exists() else []
    except (OSError, ValueError, asyncio.TimeoutError, RuntimeError) as exc:
        logger.warning("Company search spider failed: %s", exc)
        raise HTTPException(status_code=503, detail="Public web search is temporarily unavailable. Try again shortly.") from exc

    results_by_host = {}
    for company in enriched:
        if not company:
            continue
        key = (urlparse(company.get("website", "")).hostname or company.get("companyName", "").casefold()).removeprefix("www.")
        existing = results_by_host.get(key)
        if existing is None:
            results_by_host[key] = company
            continue
        for field, value in company.items():
            if field == "emails":
                existing[field] = sorted(set(existing.get(field, [])) | set(value or []))
            elif value and (not existing.get(field) or field in {"address", "description", "street", "city", "state", "postalCode", "country"}):
                existing[field] = value
        if re.search(r"contact|about|team|staff", company.get("sourceUrl", ""), re.I):
            existing["sourceUrl"] = company["sourceUrl"]

    results = list(results_by_host.values())[:10]

    payload = {
        "total": len(results),
        "results": results,
        "warnings": [] if results else ["No public web results could be opened for this search."],
    }
    COMPANY_SEARCH_CACHE[cache_key] = (time.time(), payload)
    if len(COMPANY_SEARCH_CACHE) > 100:
        oldest = min(COMPANY_SEARCH_CACHE, key=lambda key: COMPANY_SEARCH_CACHE[key][0])
        COMPANY_SEARCH_CACHE.pop(oldest, None)
    return payload


@app.post("/validate")
@app.post("/api/validate")
async def validate_phone_endpoint(request: PhoneValidationRequest):
    if not clean(request.phone):
        raise HTTPException(status_code=400, detail="Enter a phone number")
    return await check_phone(request.phone, request.country, request.useProvider)


@app.post("/validate/bulk")
@app.post("/api/validate/bulk")
async def validate_phone_bulk_endpoint(request: PhoneBulkValidationRequest):
    if len(request.numbers) > MAX_PHONE_BULK_ROWS:
        raise HTTPException(
            status_code=413,
            detail=f"Maximum {MAX_PHONE_BULK_ROWS} rows per batch",
        )

    use_provider = request.useProvider and bool(TWILIO_API_KEY and TWILIO_API_SECRET)
    semaphore = asyncio.Semaphore(3 if use_provider else 15)

    async def check_one(item):
        if isinstance(item, dict):
            phone = item.get("phone") or item.get("mobile") or item.get("number") or ""
            country = item.get("country") or request.country
        else:
            phone = item if isinstance(item, str) else ""
            country = request.country
        async with semaphore:
            return await check_phone(phone, country, use_provider)

    results = await asyncio.gather(*(check_one(item) for item in request.numbers))
    return {
        "results": results,
        "count": len(results),
        "providerConfigured": bool(TWILIO_API_KEY and TWILIO_API_SECRET),
    }


async def enrich_rows(headers, rows):
    extra_columns = ["Email Confidence", "Lookup Result"]
    output_headers = list(headers)
    for column in extra_columns:
        if column not in output_headers:
            output_headers.append(column)

    async with httpx.AsyncClient(timeout=30) as client:
        for row in rows:
            for key in headers:
                row[key] = clean(row.get(key))

            row["Email Confidence"] = ""
            row["Lookup Result"] = ""

            if row.get("Email"):
                row["Lookup Result"] = "existing_email_preserved"
                continue

            first_name = row.get("First Name", "")
            last_name = row.get("Last Name", "")

            if not first_name or not last_name:
                first_name, last_name = split_full_name(
                    row.get("Full Name", "")
                )

            domain = get_domain(row.get("Company Website", ""))

            if not first_name or not last_name or not domain:
                row["Lookup Result"] = "missing_name_or_domain"
                continue

            try:
                result = await find_with_hunter(
                    client, first_name, last_name, domain
                )

                if result["email"]:
                    row["Email"] = result["email"]
                    row["Email Status"] = result["verification"]
                    row["Email Confidence"] = result["confidence"]

                row["Lookup Result"] = result["result"]

            except (httpx.HTTPError, ValueError):
                row["Lookup Result"] = "request_failed"

    return output_headers


def read_excel(raw, extension):
    if extension == ".xlsx":
        if load_workbook is None:
            raise HTTPException(
                status_code=503,
                detail="Excel support is missing; install backend/requirements.txt",
            )
        workbook = load_workbook(io.BytesIO(raw))
        sheet = workbook.active
        values = list(sheet.iter_rows(values_only=True))
        workbook.close()
    else:
        if xlrd is None:
            raise HTTPException(
                status_code=503,
                detail="Legacy .xls support is missing; install backend/requirements.txt",
            )
        workbook = xlrd.open_workbook(file_contents=raw)
        sheet = workbook.sheet_by_index(0)
        values = [sheet.row_values(index) for index in range(sheet.nrows)]

    if not values:
        raise HTTPException(status_code=400, detail="Excel file is empty")

    headers = [clean(value) for value in values[0]]
    while headers and not headers[-1]:
        headers.pop()

    if not headers or any(not header for header in headers):
        raise HTTPException(
            status_code=400,
            detail="Excel header row contains an empty column name",
        )

    missing = REQUIRED_COLUMNS - set(headers)
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"Missing required columns: {', '.join(sorted(missing))}",
        )

    rows = []
    for values_row in values[1:]:
        values_row = list(values_row[:len(headers)])
        values_row.extend([""] * (len(headers) - len(values_row)))
        rows.append(dict(zip(headers, values_row)))

    return headers, rows


def make_xlsx(headers, rows):
    if Workbook is None:
        raise HTTPException(
            status_code=503,
            detail="Excel support is missing; install backend/requirements.txt",
        )
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


@app.post("/upload-csv")
@app.post("/upload")
@app.post("/api/upload-csv")
@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    extension = os.path.splitext(file.filename or "")[1].lower()
    if extension not in {".csv", ".xlsx", ".xls"}:
        raise HTTPException(
            status_code=400,
            detail="Please upload a CSV or Excel file (.csv, .xlsx, .xls)",
        )

    if not HUNTER_API_KEY:
        raise HTTPException(
            status_code=500,
            detail="Hunter API key is not configured",
        )

    raw = await file.read()

    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail="File must be 4 MB or smaller on Vercel",
        )

    if extension == ".csv":
        try:
            # Excel commonly exports CSV files using the Windows code page.
            # Keep UTF-8 (including BOM) as the preferred encoding, but accept
            # those files too so accented names do not prevent an upload.
            try:
                text = raw.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = raw.decode("cp1252")

            reader = csv.DictReader(io.StringIO(text))

            if not reader.fieldnames:
                raise HTTPException(status_code=400, detail="CSV is empty")

            headers = [header.strip() for header in reader.fieldnames]
            reader.fieldnames = headers

            missing = REQUIRED_COLUMNS - set(headers)
            if missing:
                raise HTTPException(
                    status_code=400,
                    detail=f"Missing required columns: {', '.join(sorted(missing))}",
                )

            rows = list(reader)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail="Could not read the CSV file",
            ) from exc
    else:
        try:
            headers, rows = read_excel(raw, extension)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail="Could not read the Excel file",
            ) from exc

    if len(rows) > 1000:
        raise HTTPException(
            status_code=400,
            detail="Maximum 1,000 rows per upload for this version",
        )

    output_headers = await enrich_rows(headers, rows)

    if extension in {".xlsx", ".xls"}:
        output_bytes = make_xlsx(output_headers, rows)
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        output_filename = "infinityinfo_enriched.xlsx"
    else:
        output = io.StringIO(newline="")
        writer = csv.DictWriter(
            output,
            fieldnames=output_headers,
            extrasaction="ignore",
            lineterminator="\n",
        )
        writer.writeheader()
        writer.writerows(rows)
        output_bytes = output.getvalue().encode("utf-8-sig")
        media_type = "text/csv; charset=utf-8"
        output_filename = "infinityinfo_enriched.csv"

    return StreamingResponse(
        io.BytesIO(output_bytes),
        media_type=media_type,
        headers={
            "Content-Disposition":
                f'attachment; filename="{output_filename}"'
        },
    )
